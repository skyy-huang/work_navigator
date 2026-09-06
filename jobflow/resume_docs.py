"""简历导出：生成可下载的 .docx / .pdf。"""

import io
import os
import re
from xml.sax.saxutils import escape
from typing import Dict, Optional


def _text(value, limit=400) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()[:limit]


def _keep_lines(value, limit=4000) -> str:
    lines = []
    for raw in str(value or "").splitlines():
        line = re.sub(r"[ \t]+", " ", raw).strip()
        if line:
            lines.append(line)
    text = "\n".join(lines)
    return text[:limit]


def _html(value) -> str:
    return escape(str(value or "")).replace("\n", "<br/>")


def _contact_lines(resume: Dict) -> str:
    parts = [_text(resume.get("phone")), _text(resume.get("email"))]
    wechat = _text(resume.get("wechat"))
    if wechat:
        parts.append("微信 " + wechat)
    return " · ".join(part for part in parts if part)


def _resume_profile(resume: Dict, profile: Optional[Dict]) -> Dict:
    profile = profile or {}
    education = (resume.get("education") or [{}])[0]
    return {
        "name": _text(resume.get("full_name"), 40) or _text(profile.get("name"), 40) or "同学",
        "school": _text(education.get("school"), 80) or _text(profile.get("school"), 80),
        "major": _text(education.get("major"), 80) or _text(profile.get("major"), 80),
        "degree": _text(education.get("degree"), 20),
        "grade": _text(profile.get("grade"), 20),
        "contact": _contact_lines(resume),
    }


def resume_to_docx(resume: Dict, profile: Optional[Dict] = None) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = int(0.7 * 914400)
        section.top_margin = section.bottom_margin = int(0.55 * 914400)
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    info = _resume_profile(resume, profile)
    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = head.add_run(info["name"])
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.name = "Microsoft YaHei"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    def add_line(text, bold=False, size=10.5):
        p = doc.add_paragraph()
        for index, segment in enumerate((text or "").split("\n")):
            if index:
                p.add_run().add_break()
            run = p.add_run(segment)
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        return p

    meta = [info.get("school"), info.get("major"), info.get("grade")]
    add_line("  ".join(v for v in meta if v), size=11)
    contact = info.get("contact")
    if contact:
        add_line(contact, size=10)

    def section(title, items):
        if not items:
            return
        add_line("", size=6)
        add_line(title, bold=True, size=14)
        for item in items:
            if isinstance(item, tuple):
                add_line(item[0], bold=True, size=11)
                body = item[1]
            else:
                body = item
            if body:
                add_line(body, size=10.5)

    objective = _text(resume.get("target_role"))
    if resume.get("target_city"):
        objective += " · " + _text(resume.get("target_city"))
    if resume.get("target_industry"):
        objective += " · " + _text(resume.get("target_industry"))
    section("求职意向", [objective] if objective else [])
    section("自我评价", [_text(resume.get("self_intro"), 800)])
    edu_items = []
    for edu in resume.get("education") or []:
        title = "  ".join(v for v in [
            _text(edu.get("school")),
            _text(edu.get("major")),
            _text(edu.get("degree")),
            _text(edu.get("period")),
        ] if v)
        edu_items.append((title, _keep_lines(edu.get("highlights"), 2000)))
    section("教育经历", edu_items)
    proj_items = []
    for proj in resume.get("projects") or []:
        title_parts = [_text(proj.get("name")), _text(proj.get("subtitle")), _text(proj.get("period"))]
        body_parts = []
        description = _keep_lines(proj.get("description"), 3000)
        achievement = _keep_lines(proj.get("achievement"), 3000)
        if description:
            body_parts.append("项目内容：" + description)
        if achievement:
            body_parts.append("项目收获：" + achievement)
        proj_items.append(("  ".join(v for v in title_parts if v), "\n".join(body_parts)))
    section("项目经历", proj_items)
    inter_items = []
    for inter in resume.get("internships") or []:
        title = "  ".join(v for v in [
            _text(inter.get("company")),
            _text(inter.get("role")),
            _text(inter.get("period")),
        ] if v)
        inter_items.append((title, _text(inter.get("description"), 1000)))
    section("实习经历", inter_items)
    if resume.get("honors"):
        section("荣誉奖项", [_keep_lines(item, 200) for item in resume.get("honors")])
    if resume.get("skills"):
        section("专业技能", [" / ".join(resume.get("skills") or [])])

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _find_cjk_font():
    candidates = [
        "C:/Windows/Fonts/simhei.ttf",
        "C:/Windows/Fonts/msyh.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Light.ttc",
    ]
    for candidate in candidates:
        if os.path.exists(candidate):
            return candidate
    return ""


def resume_to_pdf(resume: Dict, profile: Optional[Dict] = None) -> bytes:
    """用 ReportLab 输出可打印的 A4 简历（中文字体子集化后体积很小）。"""
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise RuntimeError("PDF 导出需要安装 reportlab") from exc

    font_path = _find_cjk_font()
    if font_path:
        try:
            pdfmetrics.registerFont(TTFont("ZhihangCN", font_path))
            font_name = "ZhihangCN"
        except Exception:
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            font_name = "STSong-Light"
    else:
        try:
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            font_name = "STSong-Light"
        except Exception as exc:
            raise RuntimeError("PDF 导出缺少可用的中文字体") from exc

    info = _resume_profile(resume, profile)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=44,
        rightMargin=44,
        topMargin=40,
        bottomMargin=40,
        title=f"{info['name']} - 简历",
    )
    title_style = ParagraphStyle(
        "TitleCN", fontName=font_name, fontSize=23, leading=30,
        alignment=TA_CENTER, spaceAfter=2,
    )
    meta_style = ParagraphStyle(
        "MetaCN", fontName=font_name, fontSize=10.5, leading=16,
        alignment=TA_CENTER, spaceAfter=2,
    )
    section_style = ParagraphStyle(
        "SectionCN", fontName=font_name, fontSize=13, leading=19,
        spaceBefore=12, spaceAfter=5, textColor=colors.HexColor("#1d4ed8"),
    )
    body_style = ParagraphStyle(
        "BodyCN", fontName=font_name, fontSize=10.5, leading=17,
        spaceAfter=4, wordWrap="CJK",
    )
    story = [Paragraph(_html(info["name"]), title_style)]
    meta = "　".join(v for v in [info.get("school"), info.get("major"), info.get("grade")] if v)
    if meta:
        story.append(Paragraph(_html(meta), meta_style))
    if info.get("contact"):
        story.append(Paragraph(_html(info["contact"]), meta_style))

    def add_section(title, entries):
        if not entries:
            return
        story.append(Paragraph(_html(title), section_style))
        for entry in entries:
            if isinstance(entry, tuple):
                head, body = entry
                if head:
                    story.append(Paragraph("<b>" + _html(head) + "</b>", body_style))
                if body:
                    story.append(Paragraph(_html(body), body_style))
            else:
                story.append(Paragraph(_html(entry), body_style))

    objective_parts = [v for v in [
        _text(resume.get("target_role")),
        _text(resume.get("target_city")),
        _text(resume.get("target_industry")),
    ] if v]
    add_section("求职意向", [" · ".join(objective_parts)] if objective_parts else [])
    add_section("自我评价", [_text(resume.get("self_intro"), 800)])

    edu_items = []
    for edu in resume.get("education") or []:
        title = "　".join(v for v in [
            _text(edu.get("school")), _text(edu.get("major")),
            _text(edu.get("degree")), _text(edu.get("period")),
        ] if v)
        edu_items.append((title, _keep_lines(edu.get("highlights"), 2000)))
    add_section("教育经历", edu_items)

    proj_items = []
    for proj in resume.get("projects") or []:
        title = "　".join(v for v in [
            _text(proj.get("name")), _text(proj.get("subtitle")), _text(proj.get("period")),
        ] if v)
        body_parts = []
        description = _keep_lines(proj.get("description"), 3000)
        achievement = _keep_lines(proj.get("achievement"), 3000)
        if description:
            body_parts.append("项目内容：" + description)
        if achievement:
            body_parts.append("项目收获：" + achievement)
        proj_items.append((title, "\n".join(body_parts)))
    add_section("项目经历", proj_items)

    inter_items = []
    for inter in resume.get("internships") or []:
        title = "　".join(v for v in [
            _text(inter.get("company")), _text(inter.get("role")), _text(inter.get("period")),
        ] if v)
        inter_items.append((title, _text(inter.get("description"), 1000)))
    add_section("实习经历", inter_items)
    if resume.get("honors"):
        add_section("荣誉奖项", [_keep_lines(item, 200) for item in resume.get("honors")])
    if resume.get("skills"):
        add_section("专业技能", [" / ".join(resume.get("skills") or [])])

    doc.build(story)
    return buffer.getvalue()
