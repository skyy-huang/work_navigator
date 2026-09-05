"""
RAG (Retrieval-Augmented Generation) 模块

双路检索:
  1. 图谱知识检索 — 从 master_graph.json 的节点/边/超边中语义召回
  2. 案例文档检索 — 从 /data/cases/ 的 PDF/DOCX/TXT 中语义召回

离线索引: python scripts/build_rag_index.py
运行时加载: 索引存在则自动加载，不存在则优雅降级
"""

import os
import json
from typing import List, Optional

DEFAULT_MODEL = "BAAI/bge-small-zh-v1.5"


def _get_embeddings(model_name: str = DEFAULT_MODEL, device: str = "cpu"):
    from langchain_community.embeddings import HuggingFaceEmbeddings
    return HuggingFaceEmbeddings(
        model_name=model_name,
        model_kwargs={"device": device},
        encode_kwargs={"normalize_embeddings": True},
        query_instruction="为这个句子生成表示以用于检索相关文章：",
    )


class CaseRAG:
    """
    成功案例 RAG 检索器。

    用法:
        rag = CaseRAG(graph_path, cases_dir, persist_dir)
        rag.load_index()          # 从磁盘加载已有索引
        rag.build_graph_index()   # 离线构建图谱索引
        rag.build_case_index()    # 离线构建案例索引
        rag.retrieve_graph(q)     # 运行时图谱检索
        rag.retrieve_cases(q)     # 运行时案例检索
    """

    def __init__(
        self,
        graph_path: str,
        cases_dir: str,
        persist_dir: str,
        model_name: str = DEFAULT_MODEL,
    ):
        self.graph_path = graph_path
        self.cases_dir = cases_dir
        self.persist_dir = persist_dir
        self.model_name = model_name
        self._embeddings = None
        self._graph_store = None
        self._case_store = None

    @property
    def embeddings(self):
        if self._embeddings is None:
            try:
                import torch
                device = "cuda" if torch.cuda.is_available() else "cpu"
            except Exception:
                device = "cpu"
            self._embeddings = _get_embeddings(self.model_name, device)
        return self._embeddings

    # ── 离线索引构建 ──────────────────────────────────────

    def build_graph_index(self):
        """从 master_graph.json 构建图谱向量索引"""
        from langchain.schema import Document
        from langchain_community.vectorstores import FAISS

        with open(self.graph_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        docs: List[Document] = []

        for n in data.get("nodes", []):
            desc = (n.get("description") or "").strip()
            if not desc:
                continue
            text = f"[节点] 类型:{n.get('type','?')} 名称:{n.get('id','')} 描述:{desc}"
            if n.get("industry"):
                text += f" | 行业:{n['industry']}"
            docs.append(Document(
                page_content=text,
                metadata={"kind": "node", "id": n.get("id", ""),
                           "industry": n.get("industry", ""),
                           "source_file": n.get("source_file", "")},
            ))

        for e in data.get("edges", []):
            rel = (e.get("relation") or "").strip()
            if not rel:
                continue
            text = f"[关系] {e.get('source','')} --{rel}--> {e.get('target','')}"
            if e.get("industry"):
                text += f" | 行业:{e['industry']}"
            docs.append(Document(
                page_content=text,
                metadata={"kind": "edge", "source": e.get("source", ""),
                           "target": e.get("target", ""),
                           "industry": e.get("industry", ""),
                           "source_file": e.get("source_file", "")},
            ))

        for h in data.get("hyperedges", []):
            desc = (h.get("description") or "").strip()
            nodes_str = " + ".join(h.get("nodes", []))
            text = f"[超边] {h.get('id','')} 节点组合:{nodes_str} 说明:{desc}"
            if h.get("industry"):
                text += f" | 行业:{h['industry']}"
            docs.append(Document(
                page_content=text,
                metadata={"kind": "hyperedge", "id": h.get("id", ""),
                           "industry": h.get("industry", ""),
                           "source_file": h.get("source_file", "")},
            ))

        if not docs:
            print("[RAG] 没有可索引的图谱条目")
            return

        self._graph_store = FAISS.from_documents(docs, self.embeddings)
        out = os.path.join(self.persist_dir, "graph_faiss")
        os.makedirs(out, exist_ok=True)
        self._graph_store.save_local(out)
        print(f"[RAG] 图谱索引构建完成: {len(docs)} 条 -> {out}")

    def build_case_index(self):
        """从案例文档构建文档向量索引"""
        from langchain.text_splitter import RecursiveCharacterTextSplitter
        from langchain.schema import Document
        from langchain_community.vectorstores import FAISS
        import PyPDF2
        import docx

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=800,
            chunk_overlap=200,
            separators=["\n\n", "\n", "。", "；", ".", " "],
        )
        docs: List[Document] = []
        file_count = 0

        for root, _, files in os.walk(self.cases_dir):
            for fname in files:
                ext = fname.rsplit(".", 1)[-1].lower() if "." in fname else ""
                if ext not in ("pdf", "doc", "docx", "txt"):
                    continue
                fpath = os.path.join(root, fname)
                rel = os.path.relpath(fpath, self.cases_dir).replace("\\", "/")
                parts = rel.split("/")
                industry = parts[0] if len(parts) > 1 else ""
                level = parts[1] if len(parts) > 2 else ""

                text = ""
                try:
                    if ext == "pdf":
                        with open(fpath, "rb") as f:
                            reader = PyPDF2.PdfReader(f)
                            for page in reader.pages:
                                t = page.extract_text()
                                if t:
                                    text += t + "\n"
                    elif ext in ("doc", "docx"):
                        d = docx.Document(fpath)
                        for p in d.paragraphs:
                            text += p.text + "\n"
                    elif ext == "txt":
                        with open(fpath, "r", encoding="utf-8", errors="ignore") as f:
                            text = f.read()
                except Exception as exc:
                    print(f"[RAG] 跳过 {fname}: {exc}")
                    continue

                text = text.strip()
                if len(text) < 100:
                    continue
                text = text[:30000]

                chunks = splitter.split_text(text)
                for i, chunk in enumerate(chunks):
                    docs.append(Document(
                        page_content=chunk,
                        metadata={"source_file": fname, "rel_path": rel,
                                   "industry": industry, "level": level,
                                   "chunk_idx": i},
                    ))
                file_count += 1

        if not docs:
            print("[RAG] 没有可索引的案例文档")
            return

        self._case_store = FAISS.from_documents(docs, self.embeddings)
        out = os.path.join(self.persist_dir, "cases_faiss")
        os.makedirs(out, exist_ok=True)
        self._case_store.save_local(out)
        print(f"[RAG] 案例索引构建完成: {file_count} 个文件, {len(docs)} 个片段 -> {out}")

    # ── 运行时加载 ──────────────────────────────────────

    def load_index(self):
        """从磁盘加载已有 FAISS 索引"""
        from langchain_community.vectorstores import FAISS

        gdir = os.path.join(self.persist_dir, "graph_faiss")
        if os.path.isdir(gdir):
            self._graph_store = FAISS.load_local(
                gdir, self.embeddings, allow_dangerous_deserialization=True,
            )
            print(f"[RAG] 图谱索引已加载 ({self._graph_store.index.ntotal} 条向量)")

        cdir = os.path.join(self.persist_dir, "cases_faiss")
        if os.path.isdir(cdir):
            self._case_store = FAISS.load_local(
                cdir, self.embeddings, allow_dangerous_deserialization=True,
            )
            print(f"[RAG] 案例索引已加载 ({self._case_store.index.ntotal} 条向量)")

    # ── 运行时检索 ──────────────────────────────────────

    def retrieve_graph(self, query: str, k: int = 8) -> str:
        """语义检索图谱知识，返回可直接注入 prompt 的文本"""
        if not self._graph_store:
            return ""
        results = self._graph_store.similarity_search(query, k=k)
        lines = []
        for doc in results:
            meta = doc.metadata
            line = doc.page_content
            if meta.get("source_file"):
                line += f" [来源:{meta['source_file']}]"
            lines.append(line)
        return "\n".join(lines)

    def retrieve_cases(self, query: str, k: int = 3) -> str:
        """语义检索案例文档，返回可直接注入 prompt 的文本"""
        if not self._case_store:
            return ""
        results = self._case_store.similarity_search(query, k=k)
        parts = []
        for i, doc in enumerate(results, 1):
            src = doc.metadata.get("source_file", "未知")
            ind = doc.metadata.get("industry", "")
            txt = doc.page_content[:600]
            parts.append(f"【案例片段{i}】来源:{src} 行业:{ind}\n{txt}")
        return "\n\n".join(parts)

    def is_ready(self) -> bool:
        """索引是否可用"""
        return self._graph_store is not None or self._case_store is not None
